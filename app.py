"""
CarbonPass — Streamlit app
Run locally:  streamlit run app.py
Deploy:       push to GitHub -> share.streamlit.io -> select app.py
Secrets (Streamlit Cloud -> App settings -> Secrets):
    LLM_API_KEY  = "gsk_..."
    LLM_BASE_URL = "https://api.groq.com/openai/v1"
    LLM_MODEL    = "llama-3.3-70b-versatile"
No key = demo mode (fixed sample outputs).
"""

import io
import json
import os

import pandas as pd
import streamlit as st

# Load secrets into env BEFORE importing pipeline
for k in ("LLM_API_KEY", "LLM_BASE_URL", "LLM_MODEL"):
    try:
        if k in st.secrets and not os.environ.get(k):
            os.environ[k] = st.secrets[k]
    except Exception:
        pass

import pipeline  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))

st.set_page_config(page_title="CarbonPass", page_icon="🌍", layout="wide")

st.markdown(
    """
    <style>
      .hero {padding: 1.2rem 1.4rem; border-radius: 14px;
             background: linear-gradient(135deg, #0f5132 0%, #146c43 60%, #198754 100%);
             color: white; margin-bottom: 1rem;}
      .hero h1 {color: white; margin: 0; font-size: 2rem;}
      .hero p {margin: .3rem 0 0 0; opacity: .92;}
      .draft {background:#fff3cd; color:#664d03; padding:.4rem .7rem; border-radius:8px;
              font-size:.85rem; display:inline-block; margin-bottom:.5rem;}
    </style>
    <div class="hero">
      <h1>🌍 CarbonPass</h1>
      <p>Agentic AI that gets exporters ready for the EU Carbon Border Adjustment Mechanism (CBAM):
      extract plant data → calculate embedded emissions → audit readiness → draft the EU buyer report.</p>
    </div>
    """,
    unsafe_allow_html=True,
)


# ------------------------------------------------------------ file reading
def read_upload(f):
    name = f.name.lower()
    data = f.read()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join(p.extract_text() or "" for p in reader.pages)
        if name.endswith(".docx"):
            import docx
            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)
        if name.endswith((".xlsx", ".xls")):
            sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
            return "\n\n".join(f"Sheet: {n}\n{df.to_csv(index=False)}" for n, df in sheets.items())
        if name.endswith(".csv"):
            return pd.read_csv(io.BytesIO(data)).to_csv(index=False)
        return data.decode("utf-8", errors="ignore")
    except Exception as e:
        return f"[Could not read {f.name}: {e}]"


# ------------------------------------------------------------ sidebar
llm, mode = pipeline.make_llm()
with st.sidebar:
    st.header("⚙️ Setup")
    if mode == "mock":
        st.warning("**Demo mode**: no API key set. The agents return fixed outputs for the "
                   "sample plant. Add `LLM_API_KEY` in Secrets for live AI.")
    else:
        st.success(f"**Live AI**: {mode}")

    st.header("📂 Plant documents")
    source = st.radio("Input", ["Use demo plant (Margalla Steel, fictional)", "Upload my documents"])
    uploads = []
    if source.startswith("Upload"):
        uploads = st.file_uploader(
            "Bills, production logs, purchase & meter registers, monitoring notes",
            type=["pdf", "docx", "xlsx", "xls", "csv", "txt", "md"], accept_multiple_files=True)

    st.header("🔁 What-if")
    st.caption("Actual direct emissions from your billet supplier, if you obtained them "
               "(tCO2 per tonne of billet). Default value used now: 1.9")
    supplier_see = st.slider("Supplier actual data (tCO2/t)", 0.1, 1.9, 0.9, 0.05)

    run_btn = st.button("🚀 Run CarbonPass agents", type="primary", width="stretch")

    st.divider()
    st.caption("All figures are illustrative. Real CBAM reporting must follow the EU's official "
               "methodology. Outputs are drafts for human review.")

# ------------------------------------------------------------ documents
if source.startswith("Use demo"):
    with open(os.path.join(HERE, "sample_dossier_margalla_steel.md"), encoding="utf-8") as f:
        dossier = f.read()
    # Answer key is for testing only; agents must not see it
    documents_text = dossier.split("## Expected findings")[0]
    buyer_default = dossier.split("## Document 1 — Email from EU buyer")[1].split("---")[0].strip()
else:
    documents_text = "\n\n".join(f"=== {u.name} ===\n{read_upload(u)}" for u in uploads or [])
    buyer_default = ""

with st.expander("📧 EU buyer's request", expanded=False):
    buyer_request = st.text_area("Paste the buyer's email", buyer_default, height=180)

with st.expander("📄 Documents the agents will read", expanded=False):
    st.text(documents_text[:6000] or "No documents yet.")

# ------------------------------------------------------------ run
if run_btn:
    if not documents_text.strip():
        st.error("Please upload documents or choose the demo plant.")
        st.stop()
    with st.status("Agents at work…", expanded=True) as status:
        try:
            out = pipeline.run(
                documents_text, buyer_request=buyer_request,
                supplier_whatif={"process": "Rolling Mill - Rebar",
                                 "precursor": "Billets (purchased)",
                                 "see_direct": supplier_see},
                llm=llm, on_step=lambda s: st.write("✅ " + s))
            st.session_state["out"] = out
            status.update(label="Done: report ready", state="complete")
        except Exception as e:
            status.update(label="A step failed", state="error")
            st.error(f"{type(e).__name__}: {e}")
            st.stop()

out = st.session_state.get("out")
if not out:
    st.info("👈 Choose the demo plant (or upload documents) and click **Run CarbonPass agents**.")
    st.stop()

calc = out["calculation"]
rep = out["report"]
audit = out["audit"]
exp_d = calc["eu_cost_exposure"][0] if calc["eu_cost_exposure"] else None
exp_w = out["what_if"]["eu_cost_exposure"][0] if out.get("what_if") and out["what_if"]["eu_cost_exposure"] else None
last_proc = list(calc["results"].values())[-1]

# ------------------------------------------------------------ headline
st.subheader(f"🏭 {calc['installation']}")
st.caption(f"Reporting period: {calc['reporting_period']}")

c1, c2, c3, c4 = st.columns(4)
c1.metric("Readiness score", f"{rep['readiness_score']}/100")
c2.metric("Direct emissions (final product)", f"{last_proc['SEE_direct_tCO2_per_t']} tCO2/t")
if exp_d:
    c3.metric("EU cost exposure (default values)",
              f"€{exp_d['illustrative_cost_full_phase_in_EUR']:,}")
if exp_d and exp_w:
    saving = exp_d["illustrative_cost_full_phase_in_EUR"] - exp_w["illustrative_cost_full_phase_in_EUR"]
    pct = 100 * saving / exp_d["illustrative_cost_full_phase_in_EUR"] if exp_d["illustrative_cost_full_phase_in_EUR"] else 0
    c4.metric("With actual supplier data", f"€{exp_w['illustrative_cost_full_phase_in_EUR']:,}",
              f"-€{saving:,} ({pct:.0f}%)", delta_color="inverse")

st.markdown(f'<span class="draft">DRAFT – requires review</span>', unsafe_allow_html=True)
st.write(rep.get("readiness_summary", ""))

tabs = st.tabs(["📊 Emissions", "✅ Readiness audit", "🗺️ Action plan",
                "🔁 What-if", "📧 Buyer email", "🔍 Extracted data"])

# Emissions
with tabs[0]:
    rows = []
    for name, r in calc["results"].items():
        rows.append({"Process / product": name, "CN code": r["cn_code"],
                     "Output (t)": r["output_tonnes"],
                     "Direct (tCO2/t)": r["SEE_direct_tCO2_per_t"],
                     "Indirect (tCO2/t)": r["SEE_indirect_tCO2_per_t"],
                     "Fuel CO2 (t)": r["direct_fuel_tCO2"],
                     "Process CO2 (t)": r["direct_process_tCO2"],
                     "Precursor CO2 (t)": r["precursor_direct_tCO2"]})
    st.dataframe(pd.DataFrame(rows), width="stretch", hide_index=True)
    st.markdown("**Data-quality flags**")
    for fl in calc["flags"]:
        (st.error if fl.startswith("[HIGH]") else st.warning if fl.startswith("[MEDIUM]") else st.info)(fl)
    for name, r in calc["results"].items():
        if r["precursor_basis"]:
            st.markdown(f"**Precursors for {name}**")
            st.dataframe(pd.DataFrame(r["precursor_basis"]), hide_index=True)

# Audit
with tabs[1]:
    df = pd.DataFrame(audit)
    counts = df["rating"].value_counts()
    a1, a2, a3 = st.columns(3)
    a1.metric("✅ Met", int(counts.get("Met", 0)))
    a2.metric("🟡 Partial", int(counts.get("Partial", 0)))
    a3.metric("❌ Missing", int(counts.get("Missing", 0)))
    icon = {"Met": "✅", "Partial": "🟡", "Missing": "❌"}
    df["rating"] = df["rating"].map(lambda r: f"{icon.get(r, '')} {r}")
    st.dataframe(df[["id", "area", "rating", "severity", "evidence", "fix"]],
                 width="stretch", hide_index=True)

# Action plan
with tabs[2]:
    for a in rep.get("action_plan", []):
        st.markdown(f"**{a.get('priority')}. {a.get('action')}**  \n"
                    f"Why: {a.get('why')} · Owner: {a.get('owner')} · By: {a.get('by_when')}")

# What-if
with tabs[3]:
    if exp_d and exp_w:
        chart = pd.DataFrame({
            "Scenario": ["EU default values", "Actual supplier data"],
            "Illustrative exposure (EUR)": [exp_d["illustrative_cost_full_phase_in_EUR"],
                                            exp_w["illustrative_cost_full_phase_in_EUR"]]})
        st.bar_chart(chart, x="Scenario", y="Illustrative exposure (EUR)")
        st.write(rep.get("what_if_message", ""))
        st.caption(f"Based on {exp_d['tonnes_to_eu']:,} t exported to the EU, an EU ETS price of "
                   f"€{calc['assumptions']['ets_price_eur']}/t, at full phase-in. Simplified, "
                   f"illustrative calculation. Adjust the supplier value in the sidebar and re-run.")
    else:
        st.info("No EU exports or what-if data available.")

# Buyer email
with tabs[4]:
    em = rep.get("buyer_email", {})
    st.text_input("Subject", em.get("subject", ""))
    st.text_area("Body", em.get("body", ""), height=360)

# Extracted data
with tabs[5]:
    ex = out["extracted"]
    if ex.get("_citations"):
        st.markdown("**Citations**")
        st.dataframe(pd.DataFrame(ex["_citations"]), hide_index=True, width="stretch")
    if ex.get("_missing"):
        st.markdown("**Not found in documents**")
        for m in ex["_missing"]:
            st.write("• " + m)
    st.json({k: v for k, v in ex.items() if not k.startswith("_")}, expanded=False)

# ------------------------------------------------------------ download
md = [f"# CarbonPass report — {calc['installation']}",
      f"*Reporting period: {calc['reporting_period']} · DRAFT – requires review · figures illustrative*",
      f"\n## Summary\nReadiness score: **{rep['readiness_score']}/100**\n\n{rep.get('readiness_summary', '')}",
      "\n## Embedded emissions"]
for name, r in calc["results"].items():
    md.append(f"- **{name}** (CN {r['cn_code']}): direct {r['SEE_direct_tCO2_per_t']} tCO2/t, "
              f"indirect {r['SEE_indirect_tCO2_per_t']} tCO2/t")
md.append("\n## Data-quality flags")
md += [f"- {fl}" for fl in calc["flags"]]
md.append("\n## Readiness audit\n| ID | Area | Rating | Evidence | Fix |\n|---|---|---|---|---|")
md += [f"| {a['id']} | {a['area']} | {a['rating']} | {a['evidence']} | {a.get('fix') or ''} |" for a in out["audit"]]
md.append("\n## Action plan")
md += [f"{a.get('priority')}. **{a.get('action')}** — {a.get('why')} (Owner: {a.get('owner')}, by {a.get('by_when')})"
       for a in rep.get("action_plan", [])]
md.append(f"\n## What-if\n{rep.get('what_if_message', '')}")
em = rep.get("buyer_email", {})
md.append(f"\n## Draft email to EU buyer\n**Subject:** {em.get('subject', '')}\n\n{em.get('body', '')}")

st.divider()
st.download_button("⬇️ Download full report (Markdown)", "\n".join(md),
                   file_name="CarbonPass_report.md", mime="text/markdown")
